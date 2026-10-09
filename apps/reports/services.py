"""Генерация Excel-отчётов (openpyxl).

Два входа:
  - build_period_response(start, end) — сводный файл с листами за период;
  - single_sheet_response(...) — один лист (для точечных экспортов из админки).

Данные берём из готовых сервисов (finance/analytics) и прямых read-запросов —
это read-only проекции для выгрузки, бизнес-логику не дублируем.
Decimal → float только на записи в ячейку (Excel хранит числа как float).
"""
import io
from decimal import Decimal

from django.db.models import DecimalField, F, Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_DEC = DecimalField(max_digits=20, decimal_places=2)
_HEADER_FILL = PatternFill("solid", fgColor="10B981")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _num(value):
    return float(value or 0)


def _write_sheet(ws, headers, rows, money_cols=()):
    """Заполнить лист: жирная шапка + строки + формат чисел + ширина колонок."""
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for idx in money_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=idx).number_format = "#,##0"
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 42)


def _workbook_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def single_sheet_response(filename, sheet_title, headers, rows, money_cols=()):
    """Один лист — для точечных экспортов из админки."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    _write_sheet(ws, headers, rows, money_cols)
    return _workbook_response(wb, filename)


# ---------- строки отчётов ----------

def sales_rows(start, end):
    from apps.sales.models import Sale

    # .active() — сторнированные продажи в отчёт не попадают.
    qs = Sale.objects.active().filter(
        created_at__date__gte=start, created_at__date__lte=end
    ).select_related("client").order_by("created_at")
    return [
        [s.created_at.strftime("%d.%m.%Y"), str(s.client) if s.client else "—",
         s.get_payment_type_display(), _num(s.total)]
        for s in qs
    ]


def debt_payment_rows(start, end):
    from apps.debts.models import DebtPayment

    qs = DebtPayment.objects.filter(
        created_at__date__gte=start, created_at__date__lte=end,
        is_reversed=False, debt__is_reversed=False,
    ).select_related("debt__client").order_by("created_at")
    return [
        [p.created_at.strftime("%d.%m.%Y"), str(p.debt.client), _num(p.amount), p.comment]
        for p in qs
    ]


def expense_rows(start, end):
    from apps.finance.models import CashFlow

    qs = CashFlow.objects.filter(
        direction=CashFlow.Direction.OUT, date__gte=start, date__lte=end
    ).order_by("date")
    return [
        [c.date.strftime("%d.%m.%Y"), c.get_category_display(), _num(c.amount), c.comment]
        for c in qs
    ]


def stock_rows():
    from apps.analytics.services import _products_with_stock
    from apps.warehouse.services import stock_value

    rows = []
    for p in _products_with_stock().order_by("name"):
        rows.append([p.name, p.pack_name, int(p.stock_qty or 0), _num(stock_value(p))])
    return rows


def top_product_rows(start, end):
    from apps.sales.models import SaleItem

    qs = (
        SaleItem.objects.filter(
            sale__created_at__date__gte=start, sale__created_at__date__lte=end,
            sale__is_reversed=False,
        )
        .values("product__name")
        .annotate(
            sold=Sum("packs"),
            revenue=Sum("revenue"),
        )
        .order_by("-revenue")
    )
    return [[r["product__name"], int(r["sold"] or 0), _num(r["revenue"])] for r in qs]


def debtor_rows():
    from apps.clients.models import Client

    rows = []
    for c in Client.objects.order_by("name"):
        debt = c.current_debt
        if debt > 0:
            rows.append([c.name, c.phone, _num(debt)])
    rows.sort(key=lambda x: x[2], reverse=True)
    return rows


def summary_rows(start, end):
    """Лист «Сводка».

    Прибыль здесь — по начислению (выручка − себестоимость проданного −
    расходы бизнеса), как и на дашборде: в проекте одно определение прибыли
    (см. analytics.services.month_indicators). Рядом печатаем слагаемые,
    чтобы цифру можно было проверить, не заглядывая в код.

    Закупка товара и погашение кредитов в «Расходы бизнеса» не входят:
    первая попадёт в отчёт себестоимостью в момент продажи, второе —
    уменьшение долга, а не издержка. Личные траты владельца — изъятие,
    они уменьшают кассу, но не заработок.
    """
    from apps.analytics.services import month_indicators
    from apps.debts.models import DebtPayment
    from apps.finance.services import get_cash_balance, total_expense
    from apps.sales.models import Sale

    sales_cash = Sale.objects.active().filter(
        created_at__date__gte=start, created_at__date__lte=end,
        payment_type=Sale.PaymentType.CASH,
    ).aggregate(s=Sum("total"))["s"] or Decimal("0")
    sales_debt = Sale.objects.active().filter(
        created_at__date__gte=start, created_at__date__lte=end,
        payment_type=Sale.PaymentType.DEBT,
    ).aggregate(s=Sum("total"))["s"] or Decimal("0")
    payments = DebtPayment.objects.filter(
        created_at__date__gte=start, created_at__date__lte=end,
        is_reversed=False, debt__is_reversed=False,
    ).aggregate(s=Sum("amount"))["s"] or Decimal("0")
    ind = month_indicators(start, end)

    return [
        ["Продажи — наличные", _num(sales_cash)],
        ["Продажи — реализация", _num(sales_debt)],
        ["Итого продаж", _num(sales_cash + sales_debt)],
        ["Оплаты долгов (приход)", _num(payments)],
        ["", None],
        ["Выручка", _num(ind["revenue"])],
        ["Себестоимость проданного", _num(ind["cogs"])],
        ["Расходы бизнеса", _num(ind["expenses"])],
        ["ПРИБЫЛЬ (выручка − себестоимость − расходы)", _num(ind["profit"])],
        ["", None],
        ["Все расходы деньгами (с закупкой и личными)", _num(total_expense(start, end))],
        ["Касса на текущий момент", _num(get_cash_balance())],
    ]


def build_period_response(start, end):
    """Сводный файл с листами за период [start, end]."""
    wb = Workbook()
    wb.remove(wb.active)  # убрать пустой лист по умолчанию

    _write_sheet(wb.create_sheet("Сводка"),
                 ["Показатель", "Значение, сом"], summary_rows(start, end), money_cols=(2,))
    _write_sheet(wb.create_sheet("Продажи"),
                 ["Дата", "Клиент", "Тип оплаты", "Сумма"], sales_rows(start, end), money_cols=(4,))
    _write_sheet(wb.create_sheet("Оплаты долгов"),
                 ["Дата", "Клиент", "Сумма", "Комментарий"], debt_payment_rows(start, end), money_cols=(3,))
    _write_sheet(wb.create_sheet("Расходы"),
                 ["Дата", "Категория", "Сумма", "Комментарий"], expense_rows(start, end), money_cols=(3,))
    _write_sheet(wb.create_sheet("Остатки склада"),
                 ["Товар", "Фасовка", "Остаток (фасовок)", "Стоимость"], stock_rows(), money_cols=(4,))
    _write_sheet(wb.create_sheet("ТОП товаров"),
                 ["Товар", "Продано (фасовок)", "Выручка"], top_product_rows(start, end), money_cols=(3,))
    _write_sheet(wb.create_sheet("Долги клиентов"),
                 ["Клиент", "Телефон", "Долг"], debtor_rows(), money_cols=(3,))

    filename = f"report_{start.isoformat()}_{end.isoformat()}.xlsx"
    return _workbook_response(wb, filename)
